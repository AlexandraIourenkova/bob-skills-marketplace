#!/usr/bin/env python3
"""Convert one or more CSV files into a typed, formatted .xlsx workbook.

Run:  uv run --with openpyxl python csv_to_xlsx.py data.csv [more.csv ...] --output out.xlsx

Each CSV becomes one sheet (named after the file, or use --sheet-names). Columns
are typed from their content: integers, decimals, percentages ("12.5%"), booleans
(TRUE/FALSE, yes/no), ISO dates and datetimes, and dotted or slashed dates. Anything
that is not consistent across the whole column stays text, as do values with leading
zeros (ZIP codes, article numbers) and digit strings longer than 15 characters (IDs,
phone numbers), which Excel would otherwise mangle. The header row is bold and filled
and frozen, an autofilter covers the data, and column widths follow the content.

Slashed dates (05/06/2026) are read with ONE convention per column: month/day/year by
default, day/month/year with --dayfirst. When a value such as 13/09/2026 proves that
the column is day-first (or month-first), that convention is used for the whole
column and a note is printed; when the values fit both conventions the chosen one is
reported; when they fit neither, the column stays text with a warning. Blank lines
inside the file are skipped and counted, because they shift the sheet row numbers.

Encodings: the default tries UTF-8 with or without BOM, then UTF-16 with BOM. For
exports from Windows tools pass --encoding cp1251 (Cyrillic), cp1255 (Hebrew),
cp1252 (Western Europe) or a comma-separated list of candidates to try in order.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import re
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "openpyxl is not installed. Run with: uv run --with openpyxl python csv_to_xlsx.py ...\n"
    )
    sys.exit(2)

MAX_SHEET_NAME = 31
INVALID_SHEET_CHARS = '[]:*?/\\'
INT_RE = re.compile(r"^[+-]?\d{1,15}$")
FLOAT_RE = re.compile(r"^[+-]?(\d{1,3}(,\d{3})+|\d+)?(\.\d+)?([eE][+-]?\d+)?$")
PERCENT_RE = re.compile(r"^[+-]?(\d{1,3}(,\d{3})+|\d+)?(\.\d+)?\s*%$")
LEADING_ZERO_RE = re.compile(r"^0\d+$")
LONG_DIGITS_RE = re.compile(r"^[+-]?\d{16,}$")
TRUE_WORDS = {"true", "yes", "y"}
FALSE_WORDS = {"false", "no", "n"}
DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%Y/%m/%d")
SLASHED_DMY = "%d/%m/%Y"
SLASHED_MDY = "%m/%d/%Y"
SLASHED_RE = re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")
THOUSANDS_DOT_RE = re.compile(r"^[+-]?[1-9]\d{0,2}(\.\d{3})+$")
DATETIME_FORMATS = (
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M",
    "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M",
)


class ConvertError(Exception):
    pass


# --------------------------------------------------------------------------- reading


def read_text(path: Path, encodings: list[str]) -> tuple[str, str]:
    raw = path.read_bytes()
    if not raw:
        raise ConvertError(f"{path} is empty")
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return raw.decode("utf-16"), "utf-16"
    errors = []
    for encoding in encodings:
        try:
            return raw.decode(encoding), encoding
        except (UnicodeDecodeError, LookupError) as exc:
            errors.append(f"{encoding}: {exc}")
    raise ConvertError(
        f"{path} could not be decoded with {', '.join(encodings)}. "
        "Pass --encoding cp1251 (Cyrillic), cp1255 (Hebrew), cp1252 (Western) or another codec.\n  "
        + "\n  ".join(errors)
    )


def detect_delimiter(sample: str, explicit: str | None) -> str:
    if explicit:
        return "\t" if explicit in ("\\t", "tab") else explicit
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample.splitlines() else ""
        counts = {d: first.count(d) for d in (",", ";", "\t", "|")}
        best = max(counts, key=counts.get)
        return best if counts[best] > 0 else ","


def read_csv(path: Path, encodings: list[str], delimiter: str | None, has_header: bool) -> tuple[list[str], list[list[str]], dict]:
    text, used_encoding = read_text(path, encodings)
    if text.startswith("\ufeff"):
        text = text[1:]
    sample = "\n".join(text.splitlines()[:20])
    used_delimiter = detect_delimiter(sample, delimiter)
    reader = csv.reader(io.StringIO(text), delimiter=used_delimiter)
    raw_rows = list(reader)
    is_blank = [not any(cell.strip() for cell in row) for row in raw_rows]
    rows = [row for row, blank in zip(raw_rows, is_blank) if not blank]
    if not rows:
        raise ConvertError(f"{path} has no data rows")
    last_content = max(i for i, blank in enumerate(is_blank) if not blank)
    blank_inside = sum(1 for blank in is_blank[:last_content] if blank)
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    if has_header:
        header = [h.strip() or f"Column {i + 1}" for i, h in enumerate(rows[0])]
        data = rows[1:]
    else:
        header = [f"Column {i + 1}" for i in range(width)]
        data = rows
    info = {"encoding": used_encoding, "delimiter": used_delimiter, "rows": len(data), "columns": width,
            "blank_lines": blank_inside}
    return header, data, info


# --------------------------------------------------------------------------- typing


def parse_date(text: str, slashed_format: str | None = None) -> dt.date | dt.datetime | None:
    """Parse ISO, dotted and Y/m/d dates; slashed dates only with the convention given."""
    for fmt in DATETIME_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            continue
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    if slashed_format:
        try:
            return dt.datetime.strptime(text, slashed_format).date()
        except ValueError:
            return None
    return None


def fits_format(text: str, fmt: str) -> bool:
    try:
        dt.datetime.strptime(text, fmt)
        return True
    except ValueError:
        return False


def describe_convention(fmt: str) -> str:
    return "day/month/year" if fmt == SLASHED_DMY else "month/day/year"


def slashed_convention(values: list[str], dayfirst: bool) -> tuple[str | None, str | None]:
    """Pick one slashed-date convention for a whole column. Returns (format, note)."""
    slashed = [v for v in values if SLASHED_RE.match(v)]
    if not slashed:
        return None, None
    fits_dmy = all(fits_format(v, SLASHED_DMY) for v in slashed)
    fits_mdy = all(fits_format(v, SLASHED_MDY) for v in slashed)
    preferred = SLASHED_DMY if dayfirst else SLASHED_MDY
    if fits_dmy and fits_mdy:
        hint = "" if dayfirst else "; pass --dayfirst for day/month/year"
        return preferred, (f"slashed dates fit both conventions; read as "
                           f"{describe_convention(preferred)} for the whole column{hint}")
    if fits_dmy:
        note = None if dayfirst else ("slashed dates only fit day/month/year (a day above 12 occurs); "
                                      "read as day/month/year for the whole column")
        return SLASHED_DMY, note
    if fits_mdy:
        note = None if not dayfirst else ("slashed dates only fit month/day/year (a month above 12 "
                                          "occurs); read as month/day/year despite --dayfirst")
        return SLASHED_MDY, note
    return None, ("slashed dates mix day/month/year and month/day/year values; "
                  "column kept as text, fix the source or convert the dates first")


def normalise_number(text: str, decimal_comma: bool) -> str:
    """Turn '1.500,75' / '1500,75' / '1.500' into '1500.75' / '1500.75' / '1500' when the
    file uses a decimal comma (a dot between groups of three digits is a thousands separator)."""
    if not decimal_comma:
        return text
    if "," in text:
        return text.replace(".", "").replace(",", ".")
    if THOUSANDS_DOT_RE.match(text):
        return text.replace(".", "")
    return text


class ColumnType:
    """What a column was classified as, plus the slashed-date convention and a note."""

    def __init__(self, kind: str, slashed: str | None = None, note: str | None = None):
        self.kind = kind
        self.slashed = slashed
        self.note = note


def classify_column(values: list[str], dayfirst: bool, decimal_comma: bool = False) -> ColumnType:
    """Classify as one of: empty, int, float, percent, bool, date, datetime, text."""
    stripped = [v.strip() for v in values if v is not None and v.strip() != ""]
    if not stripped:
        return ColumnType("empty")
    if any(LEADING_ZERO_RE.match(v) for v in stripped):
        return ColumnType("text")  # ZIP codes, article numbers: keep the leading zeros
    numbers = [normalise_number(v, decimal_comma) for v in stripped]
    thousands_note = None
    if decimal_comma and any(THOUSANDS_DOT_RE.match(v) for v in stripped):
        thousands_note = ("values such as '1.500' were read as thousands groups (1500) because the "
                          "file uses a decimal comma; pass --decimal-comma no if the dot is a decimal point")
    if all(INT_RE.match(v) for v in numbers):
        return ColumnType("int", note=thousands_note)
    if any(LONG_DIGITS_RE.match(v) for v in stripped):
        return ColumnType("text")  # IDs and phone numbers lose precision as numbers
    if all(FLOAT_RE.match(v) and any(ch.isdigit() for ch in v) for v in numbers):
        return ColumnType("float", note=thousands_note)
    if all(PERCENT_RE.match(v) and any(ch.isdigit() for ch in v) for v in numbers):
        return ColumnType("percent", note=thousands_note)
    if all(v.lower() in TRUE_WORDS | FALSE_WORDS for v in stripped):
        return ColumnType("bool")
    slashed, note = slashed_convention(stripped, dayfirst)
    parsed = [parse_date(v, slashed) for v in stripped]
    if all(p is not None for p in parsed):
        kind = "datetime" if any(isinstance(p, dt.datetime) for p in parsed) else "date"
        return ColumnType(kind, slashed, note)
    if slashed is None and note and all(SLASHED_RE.match(v) or parse_date(v) is not None for v in stripped):
        return ColumnType("text", None, note)  # a date column with mixed conventions
    return ColumnType("text")


def convert_value(text: str, column: ColumnType, decimal_comma: bool = False) -> Any:
    if text is None:
        return None
    value = text.strip()
    if value == "":
        return None
    kind = column.kind
    if kind == "int":
        return int(normalise_number(value, decimal_comma))
    if kind == "float":
        return float(normalise_number(value, decimal_comma).replace(",", ""))
    if kind == "percent":
        number = normalise_number(value.rstrip("%").strip(), decimal_comma)
        return float(number.replace(",", "")) / 100.0
    if kind == "bool":
        return value.lower() in TRUE_WORDS
    if kind in ("date", "datetime"):
        parsed = parse_date(value, column.slashed)
        if kind == "datetime" and isinstance(parsed, dt.date) and not isinstance(parsed, dt.datetime):
            parsed = dt.datetime(parsed.year, parsed.month, parsed.day)
        return parsed
    return text


NUMBER_FORMATS = {
    "int": "General",
    "float": "General",
    "percent": "0.0%",
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm:ss",
}


# --------------------------------------------------------------------------- writing


def sheet_name_for(path: Path, used: set[str], explicit: str | None) -> str:
    base = explicit or path.stem or "Sheet"
    cleaned = "".join("_" if ch in INVALID_SHEET_CHARS else ch for ch in base).strip("'") or "Sheet"
    name = cleaned[:MAX_SHEET_NAME]
    counter = 2
    while name.lower() in used:
        suffix = f" ({counter})"
        name = cleaned[: MAX_SHEET_NAME - len(suffix)] + suffix
        counter += 1
    used.add(name.lower())
    return name


def write_sheet(ws, header: list[str], data: list[list[str]], dayfirst: bool,
                float_format: str | None, decimal_comma: bool) -> tuple[dict[str, str], list[str]]:
    """Write header and typed rows; return ({header: kind}, [notes about typing decisions])."""
    columns: dict[int, ColumnType] = {}
    notes: list[str] = []
    for col in range(len(header)):
        columns[col] = classify_column([row[col] for row in data], dayfirst, decimal_comma)
        if columns[col].note:
            notes.append(f"column {header[col]!r}: {columns[col].note}")

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    widths = [len(h) + 2 for h in header]
    for col, text in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for r, row in enumerate(data, start=2):
        for c, text in enumerate(row, start=1):
            column = columns[c - 1]
            kind = column.kind
            value = convert_value(text, column, decimal_comma)
            cell = ws.cell(row=r, column=c, value=value)
            fmt = NUMBER_FORMATS.get(kind)
            if kind == "float" and float_format:
                fmt = float_format
            if fmt and fmt != "General":
                cell.number_format = fmt
            if value is not None:
                if isinstance(value, dt.datetime):
                    length = 19
                elif isinstance(value, dt.date):
                    length = 10
                elif isinstance(value, float) and kind == "percent":
                    length = len(f"{value * 100:.1f}%")
                elif isinstance(value, float):
                    length = len(f"{value:,.2f}")
                elif isinstance(value, str):
                    length = max(len(line) for line in value.splitlines() or [""])
                else:
                    length = len(str(value))
                widths[c - 1] = max(widths[c - 1], length)

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = max(8, min(60, width + 2))
    ws.freeze_panes = "A2"
    if data:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(data) + 1}"
    return {header[i]: columns[i].kind for i in range(len(header))}, notes


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Convert CSV files into one typed, formatted .xlsx workbook.")
    parser.add_argument("csv", nargs="+", help="one or more CSV files; each becomes a sheet")
    parser.add_argument("--output", "-o", required=True, help="output .xlsx path")
    parser.add_argument("--sheet-names", help="comma-separated sheet names, in the order of the CSV files")
    parser.add_argument("--encoding", default="utf-8-sig,utf-8",
                        help="codec or comma-separated codecs to try (default: utf-8-sig,utf-8)")
    parser.add_argument("--delimiter", help="field delimiter (default: detected; use 'tab' for TSV)")
    parser.add_argument("--no-header", action="store_true", help="the CSV has no header row")
    parser.add_argument("--dayfirst", action="store_true",
                        help="read slashed dates as day/month/year instead of month/day/year")
    parser.add_argument("--float-format", help="Excel number format for decimal columns, e.g. '#,##0.00'")
    parser.add_argument("--decimal-comma", choices=["auto", "yes", "no"], default="auto",
                        help="numbers use a decimal comma such as 1.500,75 (auto: yes when the delimiter is ';')")
    args = parser.parse_args(argv)

    output = Path(args.output)
    if output.suffix.lower() != ".xlsx":
        sys.stderr.write(f"error: --output must end with .xlsx, got {output.name!r}\n")
        return 1
    encodings = [e.strip() for e in args.encoding.split(",") if e.strip()]
    if not encodings:
        sys.stderr.write("error: --encoding must name at least one codec\n")
        return 1
    names = [n.strip() for n in args.sheet_names.split(",")] if args.sheet_names else []
    if names and len(names) != len(args.csv):
        sys.stderr.write(f"error: {len(names)} sheet names given for {len(args.csv)} CSV files\n")
        return 1

    wb = Workbook()
    used: set[str] = set()
    try:
        for index, csv_path in enumerate(args.csv):
            path = Path(csv_path)
            if not path.is_file():
                raise ConvertError(f"CSV file not found: {path}")
            header, data, info = read_csv(path, encodings, args.delimiter, not args.no_header)
            title = sheet_name_for(path, used, names[index] if names else None)
            ws = wb.active if index == 0 else wb.create_sheet()
            ws.title = title
            decimal_comma = {"yes": True, "no": False}.get(args.decimal_comma, info["delimiter"] == ";")
            kinds, notes = write_sheet(ws, header, data, args.dayfirst, args.float_format, decimal_comma)
            typed = ", ".join(f"{name}:{kind}" for name, kind in kinds.items())
            sys.stderr.write(
                f"{path.name} -> sheet {title!r}: {info['rows']} rows, {info['columns']} columns, "
                f"encoding {info['encoding']}, delimiter {info['delimiter']!r}\n  types: {typed}\n"
            )
            if info["blank_lines"]:
                sys.stderr.write(
                    f"  note: {info['blank_lines']} blank line(s) inside the file were skipped; "
                    "sheet row numbers are lower than the CSV line numbers\n"
                )
            for note in notes:
                sys.stderr.write(f"  note: {note}\n")
    except ConvertError as exc:
        sys.stderr.write(f"error: {exc}\n")
        return 1
    except (OSError, csv.Error) as exc:
        sys.stderr.write(f"error: {exc}\n")
        return 1

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(str(output.resolve()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
