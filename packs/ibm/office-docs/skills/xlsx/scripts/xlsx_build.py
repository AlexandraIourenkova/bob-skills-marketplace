#!/usr/bin/env python3
"""Build an Excel workbook (.xlsx) from a JSON specification with openpyxl.

Run:  uv run --with openpyxl python xlsx_build.py spec.json --output report.xlsx

Specification (JSON object):

{
  "output": "report.xlsx",            optional, --output on the command line wins
  "currency_symbol": "$",             optional, symbol used by the "currency" preset
  "header_fill": "D9E1F2",            optional, hex RGB of the header row fill
  "parse_iso_dates": true,            optional, turn "2026-01-31" strings into dates
  "defined_names": {"TaxRate": "Config!$B$2"},   optional workbook-level names
  "sheets": [
    {
      "name": "Sales",                required, at most 31 characters
      "headers": ["Region", "Units", "Price", "Revenue", "Date", "Share"],
      "rows": [
        ["North", 10, 12.5, "=B2*C2", "2026-01-15", 0.25],
        {"Region": "South", "Units": 4, "Price": 9, "Revenue": "=B3*C3"}
      ],
      "formats": {"Price": "currency", "Revenue": "currency",
                  "Date": "date", "Share": "percent", "F": "0.0%"},
      "totals": {"label": "Total",
                 "columns": {"Units": "SUM", "Revenue": "SUM", "Share": "AVERAGE"}},
      "widths": "auto",               or {"Region": 18, "B": 10}
      "autofilter": true,
      "freeze_header": true,
      "merges": ["H1:J1"],
      "charts": [
        {"type": "bar", "title": "Revenue by region",
         "data": ["Revenue"], "categories": "Region",
         "anchor": "H2", "x_title": "Region", "y_title": "USD"}
      ],
      "validations": [
        {"range": "A2:A200", "list": ["North", "South", "East", "West"],
         "error": "Choose a region from the list"},
        {"range": "B2:B200", "type": "whole", "operator": "greaterThanOrEqual",
         "formula1": "0", "error": "Units must be a positive whole number"}
      ]
    }
  ]
}

Rules applied while building:
  * A string cell value that starts with "=" is written as a formula. openpyxl
    stores the formula text only; Excel or LibreOffice computes the result.
  * Column keys ("formats", "totals", "widths", chart "data"/"categories") may be
    a header name (case-insensitive), an upper-case column letter that lies inside
    the sheet, or a 1-based column number. Anything else is an error, so a mistyped
    header such as "Qty" is reported instead of silently addressing column QTY.
  * A totals entry may be a function name (SUM, AVERAGE, MIN, MAX, COUNT, COUNTA,
    MEDIAN) or a formula string with the placeholders {range}, {col}, {first} and
    {last}, for example "=SUBTOTAL(9,{range})". Literal braces must be doubled.
  * The "text" preset (@) keeps ISO date strings as text instead of parsing them.
  * Format presets: currency, percent, date, datetime, integer, number, text.
    Any other string is used verbatim as an Excel number format code.
  * Percent cells expect fractions: 0.25 is shown as 25%.
  * Chart ranges may be explicit ("D1:D5", "Other!$A$1:$B$9") or column keys,
    in which case the header row plus the data rows (never the totals row) are used.
  * Validation lists longer than 255 characters are stored on a hidden "_lists"
    sheet and referenced from there, because Excel limits inline lists.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import column_index_from_string, range_boundaries
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
except ImportError:  # pragma: no cover - environment problem, not a logic path
    sys.stderr.write(
        "openpyxl is not installed. Run with: uv run --with openpyxl python xlsx_build.py ...\n"
    )
    sys.exit(2)

MAX_SHEET_NAME = 31
INVALID_SHEET_CHARS = set('[]:*?/\\')
MAX_INLINE_LIST = 255
LISTS_SHEET = "_lists"
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
ISO_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2})?$")
CELL_REF = re.compile(r"^\$?[A-Z]{1,3}\$?\d+$", re.IGNORECASE)
TOTAL_FUNCTIONS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "MEDIAN"}


class SpecError(Exception):
    """Raised for problems in the JSON specification."""


# --------------------------------------------------------------------------- helpers


def quote_sheet(name: str) -> str:
    """Quote a sheet name for use inside a formula when needed."""
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", name):
        return name
    return "'" + name.replace("'", "''") + "'"


def validate_sheet_name(name: Any) -> str:
    if not isinstance(name, str) or not name.strip():
        raise SpecError("every sheet needs a non-empty string 'name'")
    if len(name) > MAX_SHEET_NAME:
        raise SpecError(f"sheet name {name!r} is longer than {MAX_SHEET_NAME} characters")
    bad = INVALID_SHEET_CHARS.intersection(name)
    if bad:
        raise SpecError(f"sheet name {name!r} contains forbidden characters: {''.join(sorted(bad))}")
    if name.startswith("'") or name.endswith("'"):
        raise SpecError(f"sheet name {name!r} must not start or end with an apostrophe")
    return name


def format_code(preset: Any, currency_symbol: str) -> str | None:
    """Translate a preset name into an Excel number format code."""
    if preset is None:
        return None
    if not isinstance(preset, str):
        raise SpecError(f"number format must be a string, got {preset!r}")
    key = preset.strip().lower()
    presets = {
        "currency": f'{currency_symbol}#,##0.00;[Red]-{currency_symbol}#,##0.00',
        "money": f'{currency_symbol}#,##0.00;[Red]-{currency_symbol}#,##0.00',
        "percent": "0.0%",
        "percentage": "0.0%",
        "date": "yyyy-mm-dd",
        "datetime": "yyyy-mm-dd hh:mm",
        "time": "hh:mm",
        "integer": "#,##0",
        "int": "#,##0",
        "number": "#,##0.00",
        "decimal": "#,##0.00",
        "text": "@",
        "general": "General",
    }
    return presets.get(key, preset)


def is_date_code(code: str | None) -> bool:
    if not code:
        return False
    lowered = code.lower()
    return any(token in lowered for token in ("yy", "dd", "hh", "mm:ss", "h:mm"))


class SheetLayout:
    """Where things ended up on a sheet, needed by charts, totals and validations."""

    def __init__(self, name: str, headers: list[str], header_row: int, first_data_row: int,
                 last_data_row: int, totals_row: int | None, column_count: int):
        self.name = name
        self.headers = headers
        self.header_row = header_row
        self.first_data_row = first_data_row
        self.last_data_row = last_data_row
        self.totals_row = totals_row
        self.column_count = column_count

    def column_index(self, key: Any) -> int:
        """Resolve a header name, a column letter or a 1-based number to a column index.

        A bare column letter is accepted only when it is upper-case and lies inside the
        sheet (or when the sheet has no headers at all), so a mistyped header name such
        as "Qty" or "Amt" fails loudly instead of silently addressing column QTY or AMT.
        """
        if isinstance(key, bool):
            raise SpecError(f"column key {key!r} on sheet {self.name!r} is not valid")
        if isinstance(key, int):
            if key < 1:
                raise SpecError(f"column number {key} on sheet {self.name!r} must be 1 or more")
            if self.column_count and key > self.column_count:
                raise SpecError(
                    f"column number {key} on sheet {self.name!r} is beyond the last used "
                    f"column ({self.column_count})"
                )
            return key
        if not isinstance(key, str) or not key.strip():
            raise SpecError(f"column key {key!r} on sheet {self.name!r} is not valid")
        if key.strip().isdigit():
            return self.column_index(int(key.strip()))  # JSON object keys are strings
        if key in self.headers:
            return self.headers.index(key) + 1
        lowered = [h.lower() for h in self.headers]
        if key.lower() in lowered:
            return lowered.index(key.lower()) + 1
        letters = key.strip()
        if re.fullmatch(r"[A-Z]{1,3}", letters):
            try:
                index = column_index_from_string(letters)
            except ValueError:
                raise SpecError(f"column letter {letters!r} on sheet {self.name!r} is beyond XFD")
            if not self.headers or index <= self.column_count:
                return index
            raise SpecError(
                f"column {key!r} not found on sheet {self.name!r}: it is not a header and "
                f"column {letters} lies outside A..{get_column_letter(self.column_count)}; "
                f"headers are {self.headers}"
            )
        raise SpecError(
            f"column {key!r} not found on sheet {self.name!r}; headers are {self.headers} "
            "(column letters must be upper-case, for example \"C\")"
        )


# --------------------------------------------------------------------------- values


def coerce_value(value: Any, column_code: str | None, parse_iso_dates: bool, where: str) -> Any:
    """Turn a JSON value into something openpyxl can store."""
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, str):
        if value.startswith("="):
            return value  # formula, stored verbatim
        # The "text" preset (@) means: keep the string exactly as given.
        wants_date = column_code != "@" and (is_date_code(column_code) or parse_iso_dates)
        if wants_date and ISO_DATE.match(value):
            try:
                return dt.date.fromisoformat(value)
            except ValueError as exc:
                raise SpecError(f"{where}: {value!r} looks like a date but is invalid ({exc})")
        if wants_date and ISO_DATETIME.match(value):
            try:
                return dt.datetime.fromisoformat(value.replace(" ", "T"))
            except ValueError as exc:
                raise SpecError(f"{where}: {value!r} looks like a datetime but is invalid ({exc})")
        return value
    raise SpecError(f"{where}: unsupported value type {type(value).__name__} ({value!r})")


def display_width(value: Any, code: str | None) -> int:
    """Approximate number of characters a value occupies once formatted."""
    if value is None:
        return 0
    if isinstance(value, str):
        if value.startswith("="):
            return 12
        return max(len(line) for line in value.splitlines() or [""])
    if isinstance(value, bool):
        return 5
    if isinstance(value, (dt.datetime,)):
        return 16
    if isinstance(value, dt.date):
        return 10
    if isinstance(value, float):
        if code and "%" in code:
            return len(f"{value * 100:,.1f}%")
        return len(f"{value:,.2f}") + 2
    if isinstance(value, int):
        return len(f"{value:,}") + 2
    return len(str(value))


# --------------------------------------------------------------------------- sheets


def write_sheet(wb: Workbook, spec: dict, defaults: dict, first: bool) -> SheetLayout:
    name = validate_sheet_name(spec.get("name"))
    if first:
        ws = wb.active
        ws.title = name
    else:
        ws = wb.create_sheet(title=name)

    headers = spec.get("headers")
    if headers is None:
        headers = []
    if not isinstance(headers, list) or not all(isinstance(h, str) for h in headers):
        raise SpecError(f"sheet {name!r}: 'headers' must be a list of strings")
    rows = spec.get("rows", [])
    if not isinstance(rows, list):
        raise SpecError(f"sheet {name!r}: 'rows' must be a list")

    header_row = 1 if headers else 0
    first_data_row = header_row + 1
    column_count = len(headers)
    for row in rows:
        if isinstance(row, list):
            column_count = max(column_count, len(row))
    layout = SheetLayout(name, headers, header_row, first_data_row, first_data_row - 1, None,
                         column_count)

    # Resolve per-column number formats before writing so date strings can be parsed.
    formats_spec = spec.get("formats") or spec.get("number_formats") or {}
    if not isinstance(formats_spec, dict):
        raise SpecError(f"sheet {name!r}: 'formats' must be an object mapping columns to formats")
    column_codes: dict[int, str] = {}
    for key, preset in formats_spec.items():
        code = format_code(preset, defaults["currency_symbol"])
        if code:
            column_codes[layout.column_index(key)] = code

    # Header row.
    if headers:
        header_font = Font(bold=True, color=defaults["header_font_color"])
        header_fill = PatternFill("solid", fgColor=defaults["header_fill"])
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    # Data rows.
    widths: dict[int, int] = {}
    for col, text in enumerate(headers, start=1):
        widths[col] = len(text) + 2  # room for the filter button
    current_row = first_data_row
    for index, row in enumerate(rows):
        where = f"sheet {name!r} row {index + 1}"
        if isinstance(row, dict):
            if not headers:
                raise SpecError(f"{where}: object rows need 'headers' on the sheet")
            unknown = [k for k in row if k not in headers]
            if unknown:
                raise SpecError(f"{where}: keys {unknown} are not in headers {headers}")
            values = [row.get(h) for h in headers]
        elif isinstance(row, list):
            values = row
        else:
            raise SpecError(f"{where}: each row must be a list or an object")
        column_count = max(column_count, len(values))
        for col, raw in enumerate(values, start=1):
            code = column_codes.get(col)
            value = coerce_value(raw, code, defaults["parse_iso_dates"], f"{where} column {col}")
            cell = ws.cell(row=current_row, column=col, value=value)
            if code and not (isinstance(value, dt.date) and not is_date_code(code)):
                # A date keeps openpyxl's default date format when the column
                # format is not a date format; otherwise Excel shows a serial number.
                cell.number_format = code
            widths[col] = max(widths.get(col, 0), display_width(value, code))
        current_row += 1
    last_data_row = current_row - 1
    layout.last_data_row = last_data_row
    layout.column_count = column_count

    # Totals row.
    totals = spec.get("totals")
    if totals:
        if last_data_row < first_data_row:
            raise SpecError(f"sheet {name!r}: totals requested but there are no data rows")
        if isinstance(totals, list):
            totals = {"columns": {key: "SUM" for key in totals}}
        if not isinstance(totals, dict):
            raise SpecError(f"sheet {name!r}: 'totals' must be an object or a list of columns")
        totals_row = last_data_row + 1
        label = totals.get("label", "Total")
        label_col = layout.column_index(totals.get("label_column", 1))
        bold = Font(bold=True)
        top = Border(top=Side(style="thin"))
        label_cell = ws.cell(row=totals_row, column=label_col, value=label)
        label_cell.font = bold
        label_cell.border = top
        columns = totals.get("columns", {})
        if isinstance(columns, list):
            columns = {key: "SUM" for key in columns}
        if not isinstance(columns, dict):
            raise SpecError(f"sheet {name!r}: totals 'columns' must be an object or a list")
        for key, func in columns.items():
            col = layout.column_index(key)
            letter = get_column_letter(col)
            rng = f"{letter}{first_data_row}:{letter}{last_data_row}"
            if isinstance(func, str) and func.startswith("="):
                try:
                    formula = func.format(range=rng, col=letter, first=first_data_row,
                                          last=last_data_row)
                except (KeyError, IndexError, ValueError) as exc:
                    raise SpecError(
                        f"sheet {name!r}: totals formula {func!r} contains braces that are "
                        "not one of the placeholders {range}, {col}, {first}, {last}; write "
                        f"literal braces as {{{{ and }}}} ({exc})"
                    )
            else:
                fname = str(func).upper()
                if fname not in TOTAL_FUNCTIONS:
                    raise SpecError(
                        f"sheet {name!r}: totals function {func!r} is not one of "
                        f"{sorted(TOTAL_FUNCTIONS)} and is not a formula string"
                    )
                formula = f"={fname}({rng})"
            cell = ws.cell(row=totals_row, column=col, value=formula)
            cell.font = bold
            cell.border = top
            code = column_codes.get(col)
            if code:
                cell.number_format = code
        for col in range(1, max(column_count, label_col) + 1):
            ws.cell(row=totals_row, column=col).border = top
        layout.totals_row = totals_row

    # Column widths.
    widths_spec = spec.get("widths", "auto")
    if widths_spec == "auto" or widths_spec is None:
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(8, min(60, width + 2))
    elif isinstance(widths_spec, dict):
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(8, min(60, width + 2))
        for key, width in widths_spec.items():
            if not isinstance(width, (int, float)) or width <= 0:
                raise SpecError(f"sheet {name!r}: width for {key!r} must be a positive number")
            ws.column_dimensions[get_column_letter(layout.column_index(key))].width = float(width)
    else:
        raise SpecError(f"sheet {name!r}: 'widths' must be \"auto\" or an object")

    # Freeze panes and autofilter (totals row stays outside the filter range).
    if headers and spec.get("freeze_header", True):
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if headers and spec.get("autofilter", True) and last_data_row >= first_data_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(column_count)}{last_data_row}"

    # Merged ranges.
    for rng in spec.get("merges", []) or []:
        if not isinstance(rng, str):
            raise SpecError(f"sheet {name!r}: merge ranges must be strings like 'A1:D1'")
        ws.merge_cells(rng)

    return layout


# --------------------------------------------------------------------------- charts


def resolve_reference(wb: Workbook, layout: SheetLayout, target: Any, include_header: bool,
                      what: str) -> list[tuple[Reference, bool]]:
    """Return (Reference, titles_from_data) pairs for a range string or column keys."""
    ws = wb[layout.name]
    if isinstance(target, str) and (":" in target or "!" in target or CELL_REF.match(target)):
        sheet_name = layout.name
        ref = target
        if "!" in target:
            sheet_part, ref = target.rsplit("!", 1)
            sheet_name = sheet_part.strip("'").replace("''", "'")
            if sheet_name not in wb.sheetnames:
                raise SpecError(f"{what}: sheet {sheet_name!r} in range {target!r} does not exist")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
        except ValueError as exc:
            raise SpecError(f"{what}: {target!r} is not a valid range ({exc})")
        if None in (min_col, min_row, max_col, max_row):
            raise SpecError(f"{what}: {target!r} must be a bounded range such as D1:D9")
        reference = Reference(wb[sheet_name], min_col=min_col, min_row=min_row,
                              max_col=max_col, max_row=max_row)
        return [(reference, include_header)]

    keys = target if isinstance(target, list) else [target]
    if not keys:
        raise SpecError(f"{what}: no columns given")
    if layout.last_data_row < layout.first_data_row:
        raise SpecError(f"{what}: sheet {layout.name!r} has no data rows to chart")
    result = []
    for key in keys:
        col = layout.column_index(key)
        min_row = layout.header_row if (include_header and layout.header_row) else layout.first_data_row
        reference = Reference(ws, min_col=col, min_row=min_row, max_col=col,
                              max_row=layout.last_data_row)
        result.append((reference, bool(include_header and layout.header_row)))
    return result


def add_charts(wb: Workbook, layout: SheetLayout, charts: list) -> int:
    ws = wb[layout.name]
    if not isinstance(charts, list):
        raise SpecError(f"sheet {layout.name!r}: 'charts' must be a list")
    for index, spec in enumerate(charts):
        what = f"sheet {layout.name!r} chart {index + 1}"
        if not isinstance(spec, dict):
            raise SpecError(f"{what}: must be an object")
        kind = str(spec.get("type", "bar")).lower()
        if kind in ("bar", "column", "col"):
            chart = BarChart()
            chart.type = "bar" if spec.get("horizontal") else "col"
            chart.grouping = "stacked" if spec.get("stacked") else "clustered"
            if spec.get("stacked"):
                chart.overlap = 100
        elif kind == "line":
            chart = LineChart()
        elif kind == "pie":
            chart = PieChart()
        else:
            raise SpecError(f"{what}: type must be bar, line or pie, not {kind!r}")

        if "data" not in spec:
            raise SpecError(f"{what}: 'data' (range or column keys) is required")
        series = resolve_reference(wb, layout, spec["data"], spec.get("titles_from_data", True), what)
        if kind == "pie" and len(series) > 1:
            raise SpecError(f"{what}: a pie chart takes exactly one data column")
        for reference, titles in series:
            chart.add_data(reference, titles_from_data=titles)
        if "categories" in spec:
            cats = resolve_reference(wb, layout, spec["categories"], False, what)
            chart.set_categories(cats[0][0])

        chart.title = spec.get("title")
        if kind != "pie":
            if spec.get("x_title"):
                chart.x_axis.title = spec["x_title"]
            if spec.get("y_title"):
                chart.y_axis.title = spec["y_title"]
            # Newer Excel builds hide axes unless delete is explicitly false.
            chart.x_axis.delete = False
            chart.y_axis.delete = False
        chart.width = float(spec.get("width", 18))
        chart.height = float(spec.get("height", 9))
        if "style" in spec:
            chart.style = int(spec["style"])

        anchor = spec.get("anchor")
        if not anchor:
            anchor_col = get_column_letter(max(layout.column_count, 1) + 2)
            anchor = f"{anchor_col}{2 + index * 20}"
        ws.add_chart(chart, anchor)
    return len(charts)


# --------------------------------------------------------------------------- validation


def add_validations(wb: Workbook, layout: SheetLayout, validations: list) -> int:
    ws = wb[layout.name]
    if not isinstance(validations, list):
        raise SpecError(f"sheet {layout.name!r}: 'validations' must be a list")
    for index, spec in enumerate(validations):
        what = f"sheet {layout.name!r} validation {index + 1}"
        if not isinstance(spec, dict) or "range" not in spec:
            raise SpecError(f"{what}: must be an object with a 'range'")
        if "list" in spec:
            items = spec["list"]
            if not isinstance(items, list) or not items:
                raise SpecError(f"{what}: 'list' must be a non-empty list of values")
            texts = [str(item) for item in items]
            if any("," in t for t in texts):
                raise SpecError(f"{what}: list items must not contain commas; use 'source' instead")
            inline = '"' + ",".join(texts) + '"'
            if len(inline) <= MAX_INLINE_LIST:
                formula1 = inline
            else:
                formula1 = store_list(wb, texts)
            dv = DataValidation(type="list", formula1=formula1, allow_blank=spec.get("allow_blank", True))
        elif "source" in spec:
            source = str(spec["source"])
            if not source.startswith("="):
                source = "=" + source
            dv = DataValidation(type="list", formula1=source, allow_blank=spec.get("allow_blank", True))
        else:
            vtype = spec.get("type")
            if vtype not in ("whole", "decimal", "date", "time", "textLength", "custom"):
                raise SpecError(
                    f"{what}: give a 'list', a 'source' range, or a 'type' of whole, decimal, "
                    "date, time, textLength or custom"
                )
            dv = DataValidation(type=vtype, operator=spec.get("operator", "between"),
                                formula1=str(spec.get("formula1", "")),
                                formula2=str(spec["formula2"]) if "formula2" in spec else None,
                                allow_blank=spec.get("allow_blank", True))
        dv.error = spec.get("error", "The value is not allowed here")
        dv.errorTitle = spec.get("error_title", "Invalid value")
        if spec.get("prompt"):
            dv.prompt = spec["prompt"]
            dv.promptTitle = spec.get("prompt_title", "Input")
        dv.showErrorMessage = True
        dv.showInputMessage = bool(spec.get("prompt"))
        ws.add_data_validation(dv)
        dv.add(str(spec["range"]))
    return len(validations)


def store_list(wb: Workbook, texts: list[str]) -> str:
    """Write a long validation list to the hidden helper sheet and return its reference."""
    if LISTS_SHEET in wb.sheetnames:
        ws = wb[LISTS_SHEET]
        col = ws.max_column + 1 if ws.max_row > 1 or ws.cell(1, 1).value is not None else 1
    else:
        ws = wb.create_sheet(LISTS_SHEET)
        ws.sheet_state = "hidden"
        col = 1
    for row, text in enumerate(texts, start=1):
        ws.cell(row=row, column=col, value=text)
    letter = get_column_letter(col)
    return f"={quote_sheet(LISTS_SHEET)}!${letter}$1:${letter}${len(texts)}"


# --------------------------------------------------------------------------- main


def is_valid_defined_name(name: str) -> bool:
    """Excel names: letters (any script), digits, '_' and '.', not a cell or R1C1 reference."""
    if not name or len(name) > 255:
        return False
    if not (name[0].isalpha() or name[0] in "_\\"):
        return False
    if not all(ch.isalnum() or ch in "_.\\" for ch in name[1:]):
        return False
    if CELL_REF.match(name) or re.fullmatch(r"[Rr]\d*[Cc]\d*|[RrCc]", name):
        return False
    return True


def build(spec: dict, output: Path) -> dict:
    if not isinstance(spec, dict):
        raise SpecError("the specification must be a JSON object")
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise SpecError("'sheets' must be a non-empty list")
    defaults = {
        "currency_symbol": str(spec.get("currency_symbol", "$")),
        "header_fill": str(spec.get("header_fill", "D9E1F2")).lstrip("#"),
        "header_font_color": str(spec.get("header_font_color", "000000")).lstrip("#"),
        "parse_iso_dates": bool(spec.get("parse_iso_dates", True)),
    }
    wb = Workbook()
    layouts: list[SheetLayout] = []
    seen: set[str] = set()
    for index, sheet_spec in enumerate(sheets):
        if not isinstance(sheet_spec, dict):
            raise SpecError(f"sheet {index + 1}: must be an object")
        layout = write_sheet(wb, sheet_spec, defaults, first=(index == 0))
        if layout.name.lower() in seen:
            raise SpecError(f"duplicate sheet name {layout.name!r} (Excel compares names case-insensitively)")
        seen.add(layout.name.lower())
        layouts.append(layout)

    chart_count = 0
    validation_count = 0
    for sheet_spec, layout in zip(sheets, layouts):
        chart_count += add_charts(wb, layout, sheet_spec.get("charts", []) or [])
        validation_count += add_validations(wb, layout, sheet_spec.get("validations", []) or [])

    names = spec.get("defined_names", {}) or {}
    if not isinstance(names, dict):
        raise SpecError("'defined_names' must be an object mapping names to references")
    for name, ref in names.items():
        if not is_valid_defined_name(str(name)):
            raise SpecError(
                f"defined name {name!r} is not a valid Excel name (letters, digits, '_' and '.', "
                "starting with a letter or '_', and not looking like a cell reference)"
            )
        defined = DefinedName(str(name), attr_text=str(ref).lstrip("="))
        try:
            wb.defined_names[str(name)] = defined
        except TypeError:  # openpyxl < 3.1 keeps a list
            wb.defined_names.append(defined)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return {
        "sheets": [l.name for l in layouts],
        "charts": chart_count,
        "validations": validation_count,
        "defined_names": len(names),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build an .xlsx workbook from a JSON specification (see the module docstring).",
    )
    parser.add_argument("spec", help="path to the JSON specification, or '-' to read it from stdin")
    parser.add_argument("--output", "-o", help="output .xlsx path (overrides 'output' in the spec)")
    parser.add_argument("--quiet", action="store_true", help="print only the output path")
    args = parser.parse_args(argv)

    try:
        if args.spec == "-":
            spec = json.load(sys.stdin)
        else:
            spec_path = Path(args.spec)
            if not spec_path.is_file():
                raise SpecError(f"spec file not found: {spec_path}")
            with spec_path.open("r", encoding="utf-8") as handle:
                spec = json.load(handle)
    except json.JSONDecodeError as exc:
        sys.stderr.write(f"error: the spec is not valid JSON: {exc}\n")
        return 1
    except SpecError as exc:
        sys.stderr.write(f"error: {exc}\n")
        return 1

    output_value = args.output or (spec.get("output") if isinstance(spec, dict) else None)
    if not output_value:
        sys.stderr.write("error: give --output or an 'output' key in the spec\n")
        return 1
    output = Path(output_value)
    if output.suffix.lower() != ".xlsx":
        sys.stderr.write(f"error: output must end with .xlsx, got {output.name!r}\n")
        return 1

    try:
        summary = build(spec, output)
    except SpecError as exc:
        sys.stderr.write(f"error: {exc}\n")
        return 1
    except (OSError, ValueError) as exc:
        sys.stderr.write(f"error: could not build workbook: {exc}\n")
        return 1

    if not args.quiet:
        sys.stderr.write(
            f"sheets: {', '.join(summary['sheets'])}; charts: {summary['charts']}; "
            f"validations: {summary['validations']}; defined names: {summary['defined_names']}\n"
        )
    print(str(output.resolve()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
