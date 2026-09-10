#!/usr/bin/env python3
"""Describe what an Excel workbook (.xlsx) contains.

Run:  uv run --with openpyxl python xlsx_inspect.py book.xlsx [--rows 10] [--json] [--formulas]

For every sheet the report lists: dimensions, the header row, how many cells hold
formulas versus plain values, whether cached formula results are present, the first
N rows, merged cells, charts, tables, data validations, conditional formats, freeze
panes and the autofilter range. Workbook-level defined names are listed once.

Options:
  --rows N       number of data rows to preview per sheet (default 5, 0 to skip)
  --sheet NAME   inspect one sheet only (repeatable)
  --json         emit a JSON document instead of Markdown
  --formulas     list every formula cell with its cached value
  --stats        numeric column statistics through pandas (needs --with pandas);
                 only columns whose cells are stored as numbers are summarised, so
                 text codes such as "00123" never appear as numbers

Dimensions and the header row cover the used range (cells that hold a value), so a
merged title that reaches past the data or an empty sheet does not widen the report.
The script only reads the workbook; it never modifies it.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "openpyxl is not installed. Run with: uv run --with openpyxl python xlsx_inspect.py ...\n"
    )
    sys.exit(2)

ERROR_VALUES = {"#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}


def is_formula(cell: Any) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def jsonable(value: Any) -> Any:
    if isinstance(value, dt.datetime) and value.time() == dt.time(0):
        return value.date().isoformat()  # Excel stores dates as datetimes at midnight
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return str(value)
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    text = str(value)
    return text.replace("|", "\\|").replace("\n", " ")


def inspect_sheet(ws, ws_values, preview_rows: int, list_formulas: bool) -> dict:
    """Collect facts about one worksheet. ws has formulas, ws_values has cached values."""
    formulas = 0
    values = 0
    empty_cached = 0
    error_cells: list[dict] = []
    formula_cells: list[dict] = []
    last_row = 0  # used range: cells that actually hold a value
    last_col = 0

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            last_row = max(last_row, cell.row)
            last_col = max(last_col, cell.column)
            if is_formula(cell):
                formulas += 1
                cached = ws_values[cell.coordinate].value if ws_values is not None else None
                if cached is None:
                    empty_cached += 1
                elif isinstance(cached, str) and cached in ERROR_VALUES:
                    error_cells.append({"cell": cell.coordinate, "formula": cell.value, "value": cached})
                if list_formulas:
                    formula_cells.append(
                        {"cell": cell.coordinate, "formula": str(cell.value), "cached": jsonable(cached)}
                    )
            else:
                values += 1
                if isinstance(cell.value, str) and cell.value in ERROR_VALUES:
                    error_cells.append({"cell": cell.coordinate, "formula": None, "value": cell.value})

    max_row = last_row
    max_col = last_col
    header = []
    if max_row >= 1 and max_col >= 1:
        header = [jsonable(ws.cell(row=1, column=c).value) for c in range(1, max_col + 1)]
        while header and header[-1] is None:
            header.pop()  # a merged title reaching past the data is not a header cell

    preview: list[list[Any]] = []
    if preview_rows > 0 and max_row >= 2:
        last = min(max_row, 1 + preview_rows)
        for r in range(2, last + 1):
            preview.append([jsonable(ws.cell(row=r, column=c).value) for c in range(1, max_col + 1)])

    charts = getattr(ws, "_charts", [])
    chart_info = []
    for chart in charts:
        title = None
        try:
            if chart.title and chart.title.tx and chart.title.tx.rich:
                parts = []
                for paragraph in chart.title.tx.rich.p:
                    for run in paragraph.r or []:
                        parts.append(run.t)
                title = "".join(parts) or None
        except AttributeError:
            title = None
        chart_info.append({"type": type(chart).__name__, "title": title})

    validations = []
    if ws.data_validations is not None:
        for dv in ws.data_validations.dataValidation:
            validations.append({"type": dv.type, "range": str(dv.sqref), "formula1": dv.formula1})

    conditional = 0
    try:
        for cf_range in ws.conditional_formatting:
            conditional += len(cf_range.rules)
    except (TypeError, AttributeError):
        conditional = 0

    return {
        "name": ws.title,
        "state": ws.sheet_state,
        "dimensions": f"A1:{get_column_letter(max_col)}{max_row}" if max_row and max_col else "(empty)",
        "stored_dimensions": ws.dimensions,
        "max_row": max_row,
        "max_column": max_col,
        "max_column_letter": get_column_letter(max_col) if max_col else None,
        "header": header,
        "formula_cells": formulas,
        "value_cells": values,
        "formulas_without_cached_value": empty_cached,
        "error_cells": error_cells,
        "preview": preview,
        "merged_cells": [str(r) for r in ws.merged_cells.ranges],
        "charts": len(charts),
        "chart_details": chart_info,
        "tables": list(ws.tables.keys()) if hasattr(ws, "tables") else [],
        "data_validations": validations,
        "conditional_format_rules": conditional,
        "freeze_panes": ws.freeze_panes,
        "autofilter": ws.auto_filter.ref if ws.auto_filter and ws.auto_filter.ref else None,
        "formulas": formula_cells,
    }


def is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def numeric_stats(path: Path, sheet_names: list[str]) -> dict[str, Any]:
    """pandas summary of the columns whose cells are stored as numbers (--stats only).

    The sheet is read with dtype=object so pandas keeps the stored cell types: a text
    column such as "00123" stays text and is not summarised.
    """
    import pandas as pd  # pandas is allowed in the inspect script only; checked in main()

    result: dict[str, Any] = {}
    for name in sheet_names:
        try:
            frame = pd.read_excel(path, sheet_name=name, engine="openpyxl", dtype=object)
        except (ValueError, ImportError) as exc:
            result[name] = {"error": str(exc)}
            continue
        columns: dict[str, dict[str, float | None]] = {}
        for col in frame.columns:
            values = frame[col].dropna()
            if values.empty or not all(is_number(v) for v in values):
                continue
            described = pd.to_numeric(values).describe().to_dict()
            columns[str(col)] = {k: (None if v != v else float(v)) for k, v in described.items()}
        result[name] = columns
    return result


def render_markdown(report: dict, preview_rows: int, list_formulas: bool) -> str:
    lines = [f"# Workbook: {report['file']}", ""]
    lines.append(f"- Sheets: {len(report['sheets'])} ({', '.join(s['name'] for s in report['sheets'])})")
    if report["defined_names"]:
        lines.append("- Defined names:")
        for item in report["defined_names"]:
            lines.append(f"  - `{item['name']}` = `{item['value']}`")
    else:
        lines.append("- Defined names: none")
    lines.append("")

    for sheet in report["sheets"]:
        lines.append(f"## Sheet: {sheet['name']}" + (" (hidden)" if sheet["state"] != "visible" else ""))
        lines.append("")
        lines.append(f"- Dimensions: {sheet['dimensions']} ({sheet['max_row']} rows x {sheet['max_column']} columns)")
        header = ", ".join(cell_text(h) for h in sheet["header"]) if sheet["header"] else "(empty)"
        lines.append(f"- Header row: {header}")
        lines.append(f"- Formula cells: {sheet['formula_cells']}; value cells: {sheet['value_cells']}")
        if sheet["formula_cells"]:
            if sheet["formulas_without_cached_value"]:
                lines.append(
                    f"- Cached results missing for {sheet['formulas_without_cached_value']} formulas "
                    "(open and save in Excel/LibreOffice or run xlsx_recalc_check.py)"
                )
            else:
                lines.append("- Cached results: present for every formula")
        if sheet["error_cells"]:
            lines.append(f"- Error values: {len(sheet['error_cells'])}")
            for item in sheet["error_cells"][:20]:
                lines.append(f"  - {item['cell']}: {item['value']}" + (f" from `{item['formula']}`" if item["formula"] else ""))
        lines.append(f"- Merged cells: {', '.join(sheet['merged_cells']) if sheet['merged_cells'] else 'none'}")
        charts = sheet["charts"]
        if charts:
            details = ", ".join(f"{c['type']}" + (f" '{c['title']}'" if c["title"] else "") for c in sheet["chart_details"])
            lines.append(f"- Charts: {charts} ({details})")
        else:
            lines.append("- Charts: 0")
        if sheet["tables"]:
            lines.append(f"- Tables: {', '.join(sheet['tables'])}")
        if sheet["data_validations"]:
            lines.append(f"- Data validations: {len(sheet['data_validations'])}")
            for dv in sheet["data_validations"]:
                lines.append(f"  - {dv['range']}: {dv['type']} {dv['formula1'] or ''}".rstrip())
        if sheet["conditional_format_rules"]:
            lines.append(f"- Conditional format rules: {sheet['conditional_format_rules']}")
        lines.append(f"- Freeze panes: {sheet['freeze_panes'] or 'none'}; autofilter: {sheet['autofilter'] or 'none'}")
        lines.append("")

        if preview_rows > 0 and sheet["header"]:
            lines.append(f"First {min(preview_rows, len(sheet['preview']))} data rows:")
            lines.append("")
            width = max(len(sheet["header"]), max((len(r) for r in sheet["preview"]), default=0))
            cols = [f"{get_column_letter(i + 1)}: {cell_text(h)}" for i, h in enumerate(sheet["header"])]
            cols += [get_column_letter(i + 1) for i in range(len(cols), width)]
            lines.append("| row | " + " | ".join(cols) + " |")
            lines.append("|---|" + "---|" * width)
            for offset, row in enumerate(sheet["preview"], start=2):
                cells = [cell_text(v) for v in row] + [""] * (width - len(row))
                lines.append(f"| {offset} | " + " | ".join(cells) + " |")
            if not sheet["preview"]:
                lines.append("| - | " + " | ".join(["(no data rows)"] + [""] * (width - 1)) + " |")
            lines.append("")

        if list_formulas and sheet["formulas"]:
            lines.append("Formulas:")
            lines.append("")
            lines.append("| cell | formula | cached value |")
            lines.append("|---|---|---|")
            for item in sheet["formulas"]:
                lines.append(f"| {item['cell']} | `{cell_text(item['formula'])}` | {cell_text(item['cached'])} |")
            lines.append("")

    if "stats" in report:
        lines.append("## Numeric column statistics")
        lines.append("")
        stats = report["stats"]
        if "error" in stats:
            lines.append(stats["error"])
        else:
            for sheet_name, columns in stats.items():
                lines.append(f"### {sheet_name}")
                if not columns:
                    lines.append("(no numeric columns)")
                    lines.append("")
                    continue
                if "error" in columns:
                    lines.append(columns["error"])
                    lines.append("")
                    continue
                lines.append("| column | count | mean | min | max |")
                lines.append("|---|---|---|---|---|")
                for col, values in columns.items():
                    lines.append(
                        f"| {col} | {values.get('count')} | {values.get('mean')} | {values.get('min')} | {values.get('max')} |"
                    )
                lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Describe the sheets, formulas, charts and names in an .xlsx workbook.")
    parser.add_argument("path", help="workbook to inspect (.xlsx or .xlsm)")
    parser.add_argument("--rows", type=int, default=5, help="data rows to preview per sheet (default 5)")
    parser.add_argument("--sheet", action="append", help="only inspect this sheet (repeatable)")
    parser.add_argument("--json", action="store_true", help="print JSON instead of Markdown")
    parser.add_argument("--formulas", action="store_true", help="list every formula cell with its cached value")
    parser.add_argument("--stats", action="store_true", help="numeric statistics via pandas (add --with pandas)")
    args = parser.parse_args(argv)

    path = Path(args.path)
    if not path.is_file():
        sys.stderr.write(f"error: file not found: {path}\n")
        return 1
    if path.suffix.lower() not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        sys.stderr.write(f"error: {path.name} is not an OOXML workbook (.xlsx/.xlsm); convert .xls files first\n")
        return 1
    if args.rows < 0:
        sys.stderr.write("error: --rows must be 0 or more\n")
        return 1
    if args.stats:
        try:
            import pandas  # noqa: F401  (only checked here; used in numeric_stats)
        except ImportError:
            sys.stderr.write(
                "error: --stats needs pandas; run with: uv run --with openpyxl --with pandas "
                "python xlsx_inspect.py ... --stats\n"
            )
            return 2

    try:
        wb = load_workbook(path)  # formulas as text
        wb_values = load_workbook(path, data_only=True)  # cached results, if any
    except Exception as exc:  # zipfile/xml errors surface as several types
        sys.stderr.write(f"error: cannot open {path}: {exc}\n")
        return 1

    selected = wb.sheetnames
    if args.sheet:
        missing = [s for s in args.sheet if s not in wb.sheetnames]
        if missing:
            sys.stderr.write(f"error: sheet(s) not found: {missing}; available: {wb.sheetnames}\n")
            return 1
        selected = [s for s in wb.sheetnames if s in args.sheet]

    defined = []
    names = wb.defined_names
    try:
        items = names.items()  # openpyxl >= 3.1: dict-like
    except AttributeError:
        items = [(d.name, d) for d in names.definedName]  # openpyxl 3.0
    for name, definition in items:
        defined.append({"name": name, "value": definition.attr_text})
    for ws in wb.worksheets:
        try:
            for name, definition in ws.defined_names.items():
                defined.append({"name": f"{ws.title}!{name}", "value": definition.attr_text})
        except AttributeError:
            pass

    report = {
        "file": str(path),
        "sheets": [inspect_sheet(wb[name], wb_values[name], args.rows, args.formulas) for name in selected],
        "defined_names": defined,
    }
    if args.stats:
        report["stats"] = numeric_stats(path, selected)

    if args.json:
        print(json.dumps(report, indent=2, ensure_ascii=False, default=jsonable))
    else:
        sys.stdout.write(render_markdown(report, args.rows, args.formulas))
    return 0


if __name__ == "__main__":
    sys.exit(main())
