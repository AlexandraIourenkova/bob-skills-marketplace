#!/usr/bin/env python3
"""Check the formulas of an .xlsx workbook for errors.

Run:  uv run --with openpyxl python xlsx_recalc_check.py book.xlsx [--json] [--static-only]

Two layers of checking:

1. Static analysis (always runs, needs only openpyxl). Every formula is tokenized
   and the script reports:
     - unbalanced parentheses or formulas the tokenizer cannot parse
     - references to sheets that do not exist in the workbook
     - cell references outside Excel's grid (columns beyond XFD, rows beyond 1048576)
     - names that are neither a cell, a range, a defined name nor a table
     - literal #REF! tokens left behind by deleted rows or columns
     - semicolons used as argument separators (the file format requires commas)
     - functions added after Excel 2007 written without the "_xlfn." prefix
       (Excel shows #NAME? for them when a library wrote the file)
     - unknown function names (warning; typos usually end as #NAME?)
     - defined names whose target sheet does not exist or that contain #REF!
     - error values already cached in the file (#REF!, #DIV/0!, #NAME?, ...)

2. LibreOffice recalculation (when soffice is installed and --static-only is not
   given). The workbook is converted to a fresh .xlsx in a temporary directory
   with "soffice --headless --convert-to xlsx" and the user's default profile.
   LibreOffice computes every formula cell that carries no cached value, which
   is the case for every file written by openpyxl; cached values that an Excel
   session saved in manual-calculation mode are NOT refreshed by this step.
   The results are then read with data_only=True and every formula whose result
   is an Excel error is reported. --save-recalculated PATH keeps a copy of the
   recalculated workbook, which is also the easiest way to obtain cached values
   for a file written by openpyxl.

   soffice is found with shutil.which on PATH (names "soffice" and "libreoffice"),
   or taken from --soffice / the SOFFICE environment variable; an explicit value
   that does not exist is an error (exit 2), a missing soffice on PATH falls back
   to the static check with a note (or exit 2 with --require-libreoffice).

Exit codes: 0 no errors, 1 errors found, 2 LibreOffice was required but is missing
or failed (also used for usage errors).
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.formula import Tokenizer
    from openpyxl.formula.tokenizer import Token, TokenizerError
    from openpyxl.utils.cell import column_index_from_string
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "openpyxl is not installed. Run with: uv run --with openpyxl python xlsx_recalc_check.py ...\n"
    )
    sys.exit(2)

ERROR_VALUES = {"#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
MAX_COLUMNS = 16384
MAX_ROWS = 1048576

CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")
COL_RANGE_RE = re.compile(r"^\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})$")
ROW_RANGE_RE = re.compile(r"^\$?(\d+):\$?(\d+)$")


def looks_like_name(text: str) -> bool:
    """A defined-name token: a letter of any script, '_' or '\\', then letters, digits, '_', '.', '\\', '?'."""
    if not text:
        return False
    if not (text[0].isalpha() or text[0] in "_\\"):
        return False
    return all(ch.isalnum() or ch in "_.\\?" for ch in text[1:])


# Functions that exist in the Excel 2007 file-format baseline (no prefix needed).
KNOWN_FUNCTIONS = set("""
ABS ACCRINT ACCRINTM ACOS ACOSH ADDRESS AMORDEGRC AMORLINC AND AREAS ASC ASIN ASINH ATAN ATAN2
ATANH AVEDEV AVERAGE AVERAGEA AVERAGEIF AVERAGEIFS BAHTTEXT BESSELI BESSELJ BESSELK BESSELY
BETADIST BETAINV BIN2DEC BIN2HEX BIN2OCT BINOMDIST CEILING CELL CHAR CHIDIST CHIINV CHITEST
CHOOSE CLEAN CODE COLUMN COLUMNS COMBIN COMPLEX CONCATENATE CONFIDENCE CONVERT CORREL COS COSH
COUNT COUNTA COUNTBLANK COUNTIF COUNTIFS COUPDAYBS COUPDAYS COUPDAYSNC COUPNCD COUPNUM COUPPCD
COVAR CRITBINOM CUBEKPIMEMBER CUBEMEMBER CUBEMEMBERPROPERTY CUBERANKEDMEMBER CUBESET
CUBESETCOUNT CUBEVALUE CUMIPMT CUMPRINC DATE DATEVALUE DAVERAGE DAY DAYS360 DB DCOUNT DCOUNTA
DDB DEC2BIN DEC2HEX DEC2OCT DEGREES DELTA DEVSQ DGET DISC DMAX DMIN DOLLAR DOLLARDE DOLLARFR
DPRODUCT DSTDEV DSTDEVP DSUM DURATION DVAR DVARP EDATE EFFECT EOMONTH ERF ERFC ERROR.TYPE EVEN
EXACT EXP EXPONDIST FACT FACTDOUBLE FALSE FDIST FIND FINDB FINV FISHER FISHERINV FIXED FLOOR
FORECAST FREQUENCY FTEST FV FVSCHEDULE GAMMADIST GAMMAINV GAMMALN GCD GEOMEAN GESTEP
GETPIVOTDATA GROWTH HARMEAN HEX2BIN HEX2DEC HEX2OCT HLOOKUP HOUR HYPERLINK HYPGEOMDIST IF
IFERROR IMABS IMAGINARY IMARGUMENT IMCONJUGATE IMCOS IMDIV IMEXP IMLN IMLOG10 IMLOG2 IMPOWER
IMPRODUCT IMREAL IMSIN IMSQRT IMSUB IMSUM INDEX INDIRECT INFO INT INTERCEPT INTRATE IPMT IRR
ISBLANK ISERR ISERROR ISEVEN ISLOGICAL ISNA ISNONTEXT ISNUMBER ISODD ISPMT ISREF ISTEXT JIS
KURT LARGE LCM LEFT LEFTB LEN LENB LINEST LN LOG LOG10 LOGEST LOGINV LOGNORMDIST LOOKUP LOWER
MATCH MAX MAXA MDETERM MDURATION MEDIAN MID MIDB MIN MINA MINUTE MINVERSE MIRR MMULT MOD MODE
MONTH MROUND MULTINOMIAL N NA NEGBINOMDIST NETWORKDAYS NOMINAL NORMDIST NORMINV NORMSDIST
NORMSINV NOT NOW NPER NPV OCT2BIN OCT2DEC OCT2HEX ODD ODDFPRICE ODDFYIELD ODDLPRICE ODDLYIELD
OFFSET OR PEARSON PERCENTILE PERCENTRANK PERMUT PHONETIC PI PMT POISSON POWER PPMT PRICE
PRICEDISC PRICEMAT PROB PRODUCT PROPER PV QUARTILE QUOTIENT RADIANS RAND RANDBETWEEN RANK RATE
RECEIVED REPLACE REPLACEB REPT RIGHT RIGHTB ROMAN ROUND ROUNDDOWN ROUNDUP ROW ROWS RSQ RTD
SEARCH SEARCHB SECOND SERIESSUM SIGN SIN SINH SKEW SLN SLOPE SMALL SQRT SQRTPI STANDARDIZE
STDEV STDEVA STDEVP STDEVPA STEYX SUBSTITUTE SUBTOTAL SUM SUMIF SUMIFS SUMPRODUCT SUMSQ
SUMX2MY2 SUMX2PY2 SUMXMY2 SYD T TAN TANH TBILLEQ TBILLPRICE TBILLYIELD TDIST TEXT TIME
TIMEVALUE TINV TODAY TRANSPOSE TREND TRIM TRIMMEAN TRUE TRUNC TTEST TYPE UPPER VALUE VAR VARA
VARP VARPA VDB VLOOKUP WEEKDAY WEEKNUM WEIBULL WORKDAY XIRR XNPV YEAR YEARFRAC YIELD YIELDDISC
YIELDMAT ZTEST
""".split())

# Functions introduced after Excel 2007. The file format stores them as _xlfn.NAME;
# written without the prefix by a library they show #NAME? in Excel.
PREFIX_REQUIRED = set("""
AGGREGATE ARABIC ARRAYTOTEXT BASE BETA.DIST BETA.INV BINOM.DIST BINOM.DIST.RANGE BINOM.INV
BITAND BITLSHIFT BITOR BITRSHIFT BITXOR BYCOL BYROW CEILING.MATH CEILING.PRECISE CHISQ.DIST
CHISQ.DIST.RT CHISQ.INV CHISQ.INV.RT CHISQ.TEST CHOOSECOLS CHOOSEROWS COMBINA CONCAT
CONFIDENCE.NORM CONFIDENCE.T COT COTH COVARIANCE.P COVARIANCE.S CSC CSCH DAYS DECIMAL DROP
ENCODEURL ERF.PRECISE ERFC.PRECISE EXPAND EXPON.DIST F.DIST F.DIST.RT F.INV F.INV.RT F.TEST
FILTER FILTERXML FLOOR.MATH FLOOR.PRECISE FORECAST.ETS FORECAST.ETS.CONFINT
FORECAST.ETS.SEASONALITY FORECAST.ETS.STAT FORECAST.LINEAR FORMULATEXT GAMMA GAMMA.DIST
GAMMA.INV GAMMALN.PRECISE GAUSS GROUPBY HSTACK HYPGEOM.DIST IFNA IFS IMAGE IMCOSH IMCOT IMCSC
IMCSCH IMSEC IMSECH IMSINH IMTAN ISFORMULA ISO.CEILING ISOMITTED ISOWEEKNUM LAMBDA LET
LOGNORM.DIST LOGNORM.INV MAKEARRAY MAP MAXIFS MINIFS MODE.MULT MODE.SNGL MUNIT NEGBINOM.DIST
NETWORKDAYS.INTL NORM.DIST NORM.INV NORM.S.DIST NORM.S.INV NUMBERVALUE PDURATION PERCENTILE.EXC
PERCENTILE.INC PERCENTOF PERCENTRANK.EXC PERCENTRANK.INC PERMUTATIONA PHI PIVOTBY POISSON.DIST
QUARTILE.EXC QUARTILE.INC RANDARRAY RANK.AVG RANK.EQ REDUCE REGEXEXTRACT REGEXREPLACE REGEXTEST
RRI SCAN SEC SECH SEQUENCE SHEET SHEETS SKEW.P SORT SORTBY STDEV.P STDEV.S STOCKHISTORY SWITCH
T.DIST T.DIST.2T T.DIST.RT T.INV T.INV.2T T.TEST TAKE TEXTAFTER TEXTBEFORE TEXTJOIN TEXTSPLIT
TOCOL TOROW TRIMRANGE UNICHAR UNICODE UNIQUE VALUETOTEXT VAR.P VAR.S VSTACK WEBSERVICE
WEIBULL.DIST WORKDAY.INTL WRAPCOLS WRAPROWS XLOOKUP XMATCH XOR Z.TEST
""".split())

class Finding:
    def __init__(self, level: str, sheet: str, cell: str, message: str, formula: str | None = None,
                 value: Any = None, source: str = "static"):
        self.level = level  # "error" or "warning"
        self.sheet = sheet
        self.cell = cell
        self.message = message
        self.formula = formula
        self.value = value
        self.source = source

    def as_dict(self) -> dict:
        return {
            "level": self.level, "source": self.source, "sheet": self.sheet, "cell": self.cell,
            "message": self.message, "formula": self.formula, "value": self.value,
        }


# --------------------------------------------------------------------------- static


def formula_text(value: Any) -> str | None:
    if isinstance(value, str) and value.startswith("="):
        return value
    text = getattr(value, "text", None)  # ArrayFormula / DataTableFormula
    if isinstance(text, str):
        return text if text.startswith("=") else "=" + text
    return None


def unquote_sheet(part: str) -> str:
    part = part.strip()
    if part.startswith("'") and part.endswith("'") and len(part) >= 2:
        part = part[1:-1].replace("''", "'")
    return part


def split_sheet_reference(value: str) -> tuple[str | None, str]:
    """Split 'Sheet'!A1 at the first '!' outside quotes; (None, value) when there is none."""
    in_quote = False
    for index, ch in enumerate(value):
        if ch == "'":
            in_quote = not in_quote
        elif ch == "!" and not in_quote:
            return value[:index], value[index + 1:]
    return None, value


def iter_defined_names(wb) -> list[tuple[str, str, str]]:
    """Return (scope, name, target) for workbook-level and sheet-level defined names."""
    items: list[tuple[str, str, str]] = []
    try:
        pairs = list(wb.defined_names.items())  # openpyxl >= 3.1
    except AttributeError:  # openpyxl 3.0
        pairs = [(d.name, d) for d in wb.defined_names.definedName]
    for name, definition in pairs:
        items.append(("workbook", name, definition.attr_text or ""))
    for ws in wb.worksheets:
        try:
            for name, definition in ws.defined_names.items():
                items.append((ws.title, name, definition.attr_text or ""))
        except AttributeError:
            pass
    return items


def cell_sort_key(cell: str) -> tuple:
    """Sort A2 before A13 (by row, then column); non-cell labels go last, alphabetically."""
    match = CELL_RE.match(cell or "")
    if match:
        try:
            column = column_index_from_string(match.group(1).upper())
        except ValueError:
            column = 0
        return (0, int(match.group(2)), column, "")
    return (1, 0, 0, cell or "")


def check_reference(ref: str) -> str | None:
    """Return an error message when a cell/range reference lies outside the grid, else None."""
    def check_cell(text: str) -> str | None:
        match = CELL_RE.match(text)
        if not match:
            return None
        col_letters, row = match.group(1).upper(), int(match.group(2))
        try:
            col = column_index_from_string(col_letters)
        except ValueError:
            return f"column {col_letters} does not exist"
        if col > MAX_COLUMNS:
            return f"column {col_letters} is beyond XFD"
        if row < 1 or row > MAX_ROWS:
            return f"row {row} is outside 1..{MAX_ROWS}"
        return None

    if ":" in ref and not COL_RANGE_RE.match(ref) and not ROW_RANGE_RE.match(ref):
        left, right = ref.split(":", 1)
        for side in (left, right):
            if CELL_RE.match(side):
                problem = check_cell(side)
                if problem:
                    return problem
        return None
    match = COL_RANGE_RE.match(ref)
    if match:
        for letters in match.groups():
            try:
                if column_index_from_string(letters.upper()) > MAX_COLUMNS:
                    return f"column {letters.upper()} is beyond XFD"
            except ValueError:
                return f"column {letters.upper()} does not exist"
        return None
    match = ROW_RANGE_RE.match(ref)
    if match:
        for number in match.groups():
            if int(number) < 1 or int(number) > MAX_ROWS:
                return f"row {number} is outside 1..{MAX_ROWS}"
        return None
    return check_cell(ref)


def looks_like_reference(ref: str) -> bool:
    if CELL_RE.match(ref) or COL_RANGE_RE.match(ref) or ROW_RANGE_RE.match(ref):
        return True
    if ":" in ref:
        left, right = ref.split(":", 1)
        return bool(CELL_RE.match(left) and CELL_RE.match(right))
    return False


def paren_balance_raw(formula: str) -> int:
    """Depth of parentheses ignoring string literals; negative when a ')' has no '('."""
    depth = 0
    in_string = False
    for ch in formula:
        if ch == '"':
            in_string = not in_string
        elif not in_string:
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth < 0:
                    return depth
    return depth


def static_check(wb, wb_values) -> list[Finding]:
    findings: list[Finding] = []
    sheet_names = {name.lower() for name in wb.sheetnames}

    names = iter_defined_names(wb)
    defined: set[str] = {name.lower() for _scope, name, _target in names}
    tables: set[str] = set()
    for ws in wb.worksheets:
        tables.update(name.lower() for name in getattr(ws, "tables", {}).keys())

    # Defined names are formulas too: a name pointing at a missing sheet is #REF! in Excel.
    for scope, name, target in names:
        if not target.strip():
            continue
        label = f"{scope} name" if scope != "workbook" else "workbook name"
        for finding in analyse_formula("(names)", name, "=" + target.lstrip("="), sheet_names, defined, tables):
            finding.message = f"{label} {name!r} = {target}: {finding.message}"
            findings.append(finding)

    for ws in wb.worksheets:
        ws_values = wb_values[ws.title] if wb_values is not None else None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = formula_text(cell.value)
                if text is None:
                    if isinstance(cell.value, str) and cell.value in ERROR_VALUES:
                        findings.append(Finding("error", ws.title, cell.coordinate,
                                                f"cell holds the error value {cell.value}",
                                                None, cell.value))
                    continue

                if ws_values is not None:
                    cached = ws_values[cell.coordinate].value
                    if isinstance(cached, str) and cached in ERROR_VALUES:
                        findings.append(Finding("error", ws.title, cell.coordinate,
                                                f"cached result is {cached}", text, cached, "cached"))

                findings.extend(analyse_formula(ws.title, cell.coordinate, text, sheet_names, defined, tables))
    return findings


def analyse_formula(sheet: str, coord: str, text: str, sheet_names: set[str],
                    defined: set[str], tables: set[str]) -> list[Finding]:
    out: list[Finding] = []
    depth = paren_balance_raw(text)
    if depth != 0:
        out.append(Finding("error", sheet, coord, "unbalanced parentheses", text))
    try:
        tokens = Tokenizer(text).items
    except (TokenizerError, IndexError, ValueError) as exc:
        out.append(Finding("error", sheet, coord, f"formula cannot be parsed: {exc}", text))
        return out

    array_depth = 0
    for token in tokens:
        if token.type == Token.ARRAY:
            array_depth += 1 if token.subtype == Token.OPEN else -1
            continue
        if token.type == Token.SEP and token.value == ";" and array_depth <= 0:
            out.append(Finding("error", sheet, coord,
                               "semicolon used as argument separator; the file format needs commas", text))
            continue
        if token.type == Token.FUNC and token.subtype == Token.OPEN:
            name = token.value[:-1].strip()  # drop the "("
            upper = name.upper()
            if upper.startswith("_XLFN."):
                continue  # already prefixed, whatever follows is accepted
            if upper in PREFIX_REQUIRED:
                out.append(Finding("error", sheet, coord,
                                   f"{upper} must be written as _xlfn.{upper} in the file, "
                                   "otherwise Excel shows #NAME?", text))
            elif upper not in KNOWN_FUNCTIONS and not upper.startswith("_XL"):
                out.append(Finding("warning", sheet, coord,
                                   f"unknown function {upper}; a typo ends as #NAME?", text))
            continue
        if token.type == Token.OPERAND and token.subtype == Token.ERROR:
            if token.value.upper() == "#REF!":
                out.append(Finding("error", sheet, coord, "formula contains #REF! (a deleted reference)", text))
            continue
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            out.extend(analyse_range_token(sheet, coord, text, token.value, sheet_names, defined, tables))
    return out


def analyse_range_token(sheet: str, coord: str, text: str, value: str, sheet_names: set[str],
                        defined: set[str], tables: set[str]) -> list[Finding]:
    out: list[Finding] = []
    value = value.strip()
    if value.startswith("["):
        out.append(Finding("warning", sheet, coord, f"external workbook reference {value}", text))
        return out
    if "[" in value:  # structured table reference such as Table1[Amount]
        table = value.split("[", 1)[0].strip()
        if table and table.lower() not in tables:
            out.append(Finding("error", sheet, coord, f"table {table!r} does not exist", text))
        return out
    if value.lower().startswith("_xlpm."):
        return out  # LAMBDA parameter

    sheet_part, ref = split_sheet_reference(value)
    if sheet_part is not None:
        for candidate in unquote_sheet(sheet_part).split(":"):
            if candidate.lower() not in sheet_names:
                out.append(Finding("error", sheet, coord, f"sheet {candidate!r} does not exist", text))
                return out
    if ref.upper() == "#REF!":
        out.append(Finding("error", sheet, coord, "formula contains #REF! (a deleted reference)", text))
        return out
    if looks_like_reference(ref):
        problem = check_reference(ref)
        if problem:
            out.append(Finding("error", sheet, coord, f"reference {ref}: {problem}", text))
        return out
    if looks_like_name(ref):
        if ref.lower() in defined or ref.upper() in ("TRUE", "FALSE"):
            return out
        out.append(Finding("error", sheet, coord, f"unknown name {ref!r} (Excel shows #NAME?)", text))
        return out
    out.append(Finding("warning", sheet, coord, f"unrecognised reference {value!r}", text))
    return out


# --------------------------------------------------------------------------- libreoffice


def find_soffice(explicit: str | None) -> tuple[str | None, str | None]:
    """Locate the LibreOffice executable with shutil.which. Returns (path, error).

    --soffice and the SOFFICE environment variable are taken literally (a path or a
    command name) and must resolve, otherwise an error message is returned. Without
    them the names "soffice" and "libreoffice" are searched on PATH; (None, None)
    means LibreOffice is simply not installed.
    """
    for label, candidate in (("--soffice", explicit), ("SOFFICE", os.environ.get("SOFFICE"))):
        if not candidate:
            continue
        found = shutil.which(candidate)
        if found:
            return found, None
        return None, f"{label} names {candidate!r}, which is not an executable file or a command on PATH"
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found, None
    return None, None


def recalculate_with_libreoffice(soffice: str, source: Path, workdir: Path, timeout: int) -> Path:
    """Convert the workbook with LibreOffice (default profile); return the recalculated copy.

    Formula cells without a cached value are computed during the conversion. Cached
    values already present in the file are kept as they are.
    """
    outdir = workdir / "out"
    outdir.mkdir(exist_ok=True)
    # Copy under a neutral name so the output name is predictable.
    staged = workdir / ("input" + source.suffix.lower())
    shutil.copyfile(source, staged)
    command = [
        soffice,
        "--headless", "--norestore", "--nologo", "--nodefault",
        "--convert-to", "xlsx", "--outdir", str(outdir), str(staged),
    ]
    try:
        completed = subprocess.run(command, capture_output=True, text=True, timeout=timeout, check=False)
    except subprocess.TimeoutExpired:
        raise RuntimeError(f"LibreOffice did not finish within {timeout} seconds")
    except OSError as exc:
        raise RuntimeError(f"could not start LibreOffice: {exc}")
    result = outdir / "input.xlsx"
    if completed.returncode != 0 or not result.is_file():
        detail = (completed.stderr or completed.stdout or "").strip()
        raise RuntimeError(f"LibreOffice conversion failed (exit {completed.returncode}): {detail[:500]}")
    return result


def libreoffice_check(wb, recalculated: Path) -> tuple[list[Finding], dict]:
    findings: list[Finding] = []
    try:
        wb_recalc = load_workbook(recalculated, data_only=True)
    except Exception as exc:
        raise RuntimeError(f"cannot read the recalculated workbook: {exc}")
    computed = 0
    missing = 0
    for ws in wb.worksheets:
        if ws.title not in wb_recalc.sheetnames:
            findings.append(Finding("warning", ws.title, "-", "sheet missing after LibreOffice conversion",
                                    source="libreoffice"))
            continue
        ws_recalc = wb_recalc[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                text = formula_text(cell.value)
                if text is None:
                    continue
                value = ws_recalc[cell.coordinate].value
                if value is None:
                    missing += 1
                    continue
                computed += 1
                if isinstance(value, str) and value in ERROR_VALUES:
                    findings.append(Finding("error", ws.title, cell.coordinate,
                                            f"recalculated result is {value}", text, value, "libreoffice"))
    return findings, {"computed": computed, "without_result": missing}


# --------------------------------------------------------------------------- output


def render_text(report: dict) -> str:
    lines = [f"# Formula check: {report['file']}", ""]
    lines.append(f"- Formulas: {report['formula_count']}")
    lines.append(f"- Static analysis: {report['static']['errors']} errors, {report['static']['warnings']} warnings")
    lo = report["libreoffice"]
    if lo["status"] == "done":
        lines.append(
            f"- LibreOffice recalculation: {lo['errors']} error cells; {lo['computed']} formulas computed, "
            f"{lo['without_result']} without a result ({lo['path']})"
        )
    else:
        lines.append(f"- LibreOffice recalculation: {lo['status']}" + (f" ({lo['detail']})" if lo.get("detail") else ""))
    if report.get("saved_copy"):
        lines.append(f"- Recalculated copy saved to: {report['saved_copy']}")
    elif report.get("save_requested"):
        lines.append(f"- WARNING: no recalculated copy was written to {report['save_requested']} "
                     "(LibreOffice did not run)")
    lines.append("")
    findings = report["findings"]
    if not findings:
        lines.append("No problems found.")
    else:
        lines.append("| level | source | sheet | cell | problem | formula |")
        lines.append("|---|---|---|---|---|---|")
        for item in findings:
            formula = (item["formula"] or "").replace("|", "\\|")
            message = item["message"].replace("|", "\\|")
            lines.append(
                f"| {item['level']} | {item['source']} | {item['sheet']} | {item['cell']} | "
                f"{message} | `{formula}` |"
            )
    lines.append("")
    lines.append(f"Result: {report['result']}")
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Find broken formulas in an .xlsx workbook (static analysis plus optional LibreOffice recalculation).",
    )
    parser.add_argument("path", help="workbook to check (.xlsx or .xlsm)")
    parser.add_argument("--json", action="store_true", help="print a JSON report instead of text")
    parser.add_argument("--static-only", action="store_true", help="skip LibreOffice even when it is installed")
    parser.add_argument("--require-libreoffice", action="store_true",
                        help="exit 2 instead of falling back to the static check when soffice is missing")
    parser.add_argument("--soffice", help="path or command name of the soffice executable (default: search PATH)")
    parser.add_argument("--timeout", type=int, default=180, help="seconds to wait for LibreOffice (default 180)")
    parser.add_argument("--save-recalculated", metavar="PATH",
                        help="also write the LibreOffice-recalculated workbook (with cached values) to PATH")
    parser.add_argument("--strict", action="store_true", help="treat warnings as errors for the exit code")
    args = parser.parse_args(argv)

    path = Path(args.path)
    if not path.is_file():
        sys.stderr.write(f"error: file not found: {path}\n")
        return 2
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        sys.stderr.write(f"error: {path.name} is not an .xlsx/.xlsm workbook\n")
        return 2
    if args.save_recalculated and not args.save_recalculated.lower().endswith(".xlsx"):
        sys.stderr.write("error: --save-recalculated must end with .xlsx\n")
        return 2

    try:
        wb = load_workbook(path)
        wb_values = load_workbook(path, data_only=True)
    except Exception as exc:
        sys.stderr.write(f"error: cannot open {path}: {exc}\n")
        return 2

    formula_count = sum(
        1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row
        if formula_text(cell.value) is not None
    )
    findings = static_check(wb, wb_values)

    lo_report: dict[str, Any] = {"status": "skipped", "errors": 0, "computed": 0, "without_result": 0}
    saved_copy = None
    if args.static_only:
        lo_report["status"] = "skipped (--static-only)"
    else:
        soffice, problem = find_soffice(args.soffice)
        if problem:
            sys.stderr.write(f"error: {problem}\n")
            return 2
        if not soffice:
            if args.require_libreoffice:
                sys.stderr.write(
                    "error: LibreOffice (soffice) was not found on PATH; install it or drop --require-libreoffice\n"
                )
                return 2
            lo_report["status"] = "not available"
            lo_report["detail"] = "soffice not found on PATH; static analysis only"
        else:
            with tempfile.TemporaryDirectory(prefix="xlsx_recalc_") as tmp:
                try:
                    recalculated = recalculate_with_libreoffice(soffice, path, Path(tmp), args.timeout)
                    lo_findings, counts = libreoffice_check(wb, recalculated)
                    if args.save_recalculated:
                        target = Path(args.save_recalculated)
                        target.parent.mkdir(parents=True, exist_ok=True)
                        shutil.copyfile(recalculated, target)
                        saved_copy = str(target.resolve())
                except RuntimeError as exc:
                    if args.require_libreoffice:
                        sys.stderr.write(f"error: {exc}\n")
                        return 2
                    lo_report["status"] = "failed"
                    lo_report["detail"] = str(exc)
                else:
                    findings.extend(lo_findings)
                    lo_report.update({
                        "status": "done", "path": soffice,
                        "errors": sum(1 for f in lo_findings if f.level == "error"),
                        **counts,
                    })

    if args.save_recalculated and saved_copy is None:
        sys.stderr.write(
            f"warning: no recalculated copy was written to {args.save_recalculated} "
            f"(LibreOffice recalculation: {lo_report['status']})\n"
        )

    static_errors = sum(1 for f in findings if f.source in ("static", "cached") and f.level == "error")
    static_warnings = sum(1 for f in findings if f.source in ("static", "cached") and f.level == "warning")
    total_errors = sum(1 for f in findings if f.level == "error")
    total_warnings = sum(1 for f in findings if f.level == "warning")
    failed = total_errors > 0 or (args.strict and total_warnings > 0)

    order = {"error": 0, "warning": 1}
    findings.sort(key=lambda f: (order.get(f.level, 2), f.sheet, cell_sort_key(f.cell)))
    report = {
        "file": str(path),
        "formula_count": formula_count,
        "static": {"errors": static_errors, "warnings": static_warnings},
        "libreoffice": lo_report,
        "saved_copy": saved_copy,
        "save_requested": args.save_recalculated if saved_copy is None else None,
        "findings": [f.as_dict() for f in findings],
        "result": "FAIL" if failed else "OK",
    }
    if args.json:
        print(json.dumps(report, indent=2, ensure_ascii=False, default=str))
    else:
        sys.stdout.write(render_text(report))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
